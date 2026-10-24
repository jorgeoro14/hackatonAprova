"""Verify supported office formats without any external documents."""

import pytest
from agents.app.rag.loaders import load_documents
from agents.app.rag.schemas import CorpusManifest

pytest.importorskip("pypdf")
pytest.importorskip("docx")


def manifest(path):
    return CorpusManifest.model_validate(
        {
            "scope": {
                "municipality_id": "example",
                "municipality_name": "Example City",
                "corpus_id": "test",
            },
            "data_kind": "synthetic",
            "documents": [
                {
                    "document_id": "sample",
                    "title": "Synthetic sample",
                    "path": path,
                    "kind": "procedure",
                }
            ],
        }
    )


def test_docx_keeps_paragraph_and_table_text(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("Synthetic review procedure")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Responsible role"
    table.cell(0, 1).text = "Fictional reviewer"
    document.add_paragraph("This is invented test material.")
    document.save(tmp_path / "sample.docx")
    pages = load_documents(manifest("sample.docx"), tmp_path)
    text = pages[0].page_content
    assert text.index("procedure") < text.index("Responsible role") < text.index("invented")
    assert "Fictional reviewer" in text
    assert pages[0].metadata["page"] == 0


def test_blank_pdf_requires_text_extraction_or_ocr(tmp_path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(tmp_path / "sample.pdf")
    with pytest.raises(ValueError, match="OCR"):
        load_documents(manifest("sample.pdf"), tmp_path)


def test_unsupported_format_is_explicitly_rejected(tmp_path):
    (tmp_path / "sample.bin").write_bytes(b"Synthetic bytes")
    with pytest.raises(ValueError, match="Supported"):
        load_documents(manifest("sample.bin"), tmp_path)


pytestmark = pytest.mark.phase0
